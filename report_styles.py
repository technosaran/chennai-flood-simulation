import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def setup_page_layout(section, left=1.25, right=1.0, top=1.0, bottom=1.0):
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69) # A4
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)

def add_page_number_to_section(section, start_at=None, num_format='decimal'):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    # Remove any existing runs completely from paragraph
    for r in list(p.runs):
        p._p.remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:fmt'), num_format)
    if start_at is not None:
        pgNumType.set(qn('w:start'), str(start_at))

def add_p(doc, text="", font_size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, 
          space_before=0, space_after=6, line_spacing=1.5, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color
    return p

def add_h1(doc, text):
    """Chapter Title: 16pt Bold Uppercase, Centered"""
    return add_p(doc, text.upper(), font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=12, line_spacing=1.5)

def add_h2(doc, text):
    """Major Heading: 14pt Bold, Left"""
    return add_p(doc, text, font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=6, line_spacing=1.5)

def add_h3(doc, text):
    """Sub Heading: 13pt Bold Italic, Left"""
    return add_p(doc, text, font_size=13, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4, line_spacing=1.5)

def add_body(doc, text):
    """Body Text: 12pt Times New Roman, Justified, 1.5 line spacing, 6pt after"""
    return add_p(doc, text, font_size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.5)

def add_bullet(doc, title, desc=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    
    r1 = p.add_run(title + (": " if desc else ""))
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    if desc:
        r2 = p.add_run(desc)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
        r2.font.bold = False
    return p

def add_fig(doc, img_path, caption, width=Inches(5.6), max_h=None):
    if os.path.exists(img_path):
        from PIL import Image
        im = Image.open(img_path)
        w_px, h_px = im.size
        aspect = w_px / h_px
        
        target_w = width
        target_h = target_w / aspect
        if max_h and target_h > max_h:
            target_h = max_h
            target_w = target_h * aspect
            
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run = p_img.add_run()
        run.add_picture(img_path, width=int(target_w), height=int(target_h))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        p_cap.paragraph_format.keep_with_next = True
        r = p_cap.add_run(caption)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.italic = True
    else:
        print(f"Warning: Image file not found: {img_path}")

def style_academic_table(table, col_widths, headers, data, alignments=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(title)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
        
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.add_row()
        fill_hex = "F2F4F8" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            p = cell.paragraphs[0]
            align = alignments[c_idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            p.alignment = align
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9.5)
            if fill_hex != "FFFFFF":
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
                cell._tc.get_or_add_tcPr().append(shd)

def add_code_block(doc, title, code_lines):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r_t = p_title.add_run(title)
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F9"/>')
    p._p.get_or_add_pPr().append(shd)
    
    # Border
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="1B365D"/></w:pBdr>')
    p._p.get_or_add_pPr().append(pBdr)
    
    run = p.add_run("\n".join(code_lines))
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
