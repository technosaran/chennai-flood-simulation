import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_styles import add_p, add_body, add_h1, add_h2, add_h3, style_academic_table
import os

def generate_front_matter(doc):
    # ==========================================
    # 1. COVER PAGE & TITLE PAGE (APPENDIX 1)
    # ==========================================
    add_p(doc, "AI-DRIVEN CRISIS RESOURCE INTELLIGENCE &\nAUTONOMOUS LOGISTICS PLATFORM", 
          font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=14)
    
    add_p(doc, "A PROJECT REPORT", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=10)
    add_p(doc, "Submitted by", font_size=14, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=8)
    
    add_p(doc, "SAJITH AR [211423247183]\nSARAN R [211423247189]", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=16)
    
    add_p(doc, "in partial fulfillment for the award of the degree\nof", font_size=14, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=6)
    add_p(doc, "BACHELOR OF TECHNOLOGY", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=4)
    add_p(doc, "in", font_size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=4)
    add_p(doc, "ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=16)
    
    # College crest
    crest_path = r"extracted_review0_assets\slide_1_img_1.png"
    if os.path.exists(crest_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(8)
        p_logo.paragraph_format.space_after = Pt(12)
        p_logo.add_run().add_picture(crest_path, width=Inches(1.85))
    
    add_p(doc, "PANIMALAR ENGINEERING COLLEGE", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=4)
    add_p(doc, "(An Autonomous Institution, Affiliated to Anna University, Chennai)", font_size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=10)
    add_p(doc, "OCTOBER 2026", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=12)
    doc.add_page_break()

    # ==========================================
    # 2. BONAFIDE CERTIFICATE (APPENDIX 2)
    # ==========================================
    add_p(doc, "PANIMALAR ENGINEERING COLLEGE", font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=4)
    add_p(doc, "(An Autonomous Institution, Affiliated to Anna University, Chennai)", font_size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=16)
    add_p(doc, "BONAFIDE CERTIFICATE", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=18)
    
    add_body(doc, "Certified that this project report titled \"AI-DRIVEN CRISIS RESOURCE INTELLIGENCE & AUTONOMOUS LOGISTICS PLATFORM\" is the bonafide work of SAJITH AR (211423247183) & SARAN R (211423247189), who carried out the project work under my supervision.")
    
    add_p(doc, "", space_before=18)
    
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    sig_table.columns[0].width = Inches(3.2)
    sig_table.columns[1].width = Inches(3.2)
    
    c_left = sig_table.rows[0].cells[0]
    c_right = sig_table.rows[0].cells[1]
    
    p_hod = c_left.paragraphs[0]
    p_hod.paragraph_format.line_spacing = 1.25
    r = p_hod.add_run("SIGNATURE\n\n\n\n")
    r.font.bold = True
    p_hod.add_run("Dr. G. SURESH M.Tech., Ph.D.,\nHead of the Department,\nProfessor,\nDepartment of Artificial Intelligence\nand Machine Learning,\nPanimalar Engineering College,\nChennai – 600 123")
    for r in p_hod.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        
    p_sup = c_right.paragraphs[0]
    p_sup.paragraph_format.line_spacing = 1.25
    r = p_sup.add_run("SIGNATURE\n\n\n\n")
    r.font.bold = True
    p_sup.add_run("Mrs. UDHAYA J.M., M.E.,\nSUPERVISOR,\nAssistant Professor,\nDepartment of Artificial Intelligence\nand Machine Learning,\nPanimalar Engineering College,\nChennai – 600 123")
    for r in p_sup.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        
    add_p(doc, "", space_before=24)
    add_p(doc, "Certified that the above candidate(s) were examined in the End Semester Project Viva-Voce Examination held on ............................................", 
          font_size=12, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=18, space_after=24)
    
    viva_table = doc.add_table(rows=1, cols=2)
    viva_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    viva_table.columns[0].width = Inches(3.2)
    viva_table.columns[1].width = Inches(3.2)
    v1 = viva_table.rows[0].cells[0].paragraphs[0]
    v1.add_run("INTERNAL EXAMINER").font.bold = True
    v2 = viva_table.rows[0].cells[1].paragraphs[0]
    v2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    v2.add_run("EXTERNAL EXAMINER").font.bold = True
    for c in [v1, v2]:
        for r in c.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            
    doc.add_page_break()

    # ==========================================
    # 3. DECLARATION BY THE STUDENT
    # ==========================================
    add_p(doc, "DECLARATION BY THE STUDENT", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18)
    add_body(doc, "WE SAJITH AR (211423247183) & SARAN R (211423247189), HEREBY DECLARE THAT THIS PROJECT REPORT TITLED \"AI-DRIVEN CRISIS RESOURCE INTELLIGENCE & AUTONOMOUS LOGISTICS PLATFORM\" UNDER THE GUIDANCE OF Mrs. UDHAYA J.M., M.E., ASSISTANT PROFESSOR IS THE ORIGINAL WORK DONE BY US, AND WE HAVE NOT PLAGIARIZED OR SUBMITTED TO ANY OTHER DEGREE IN ANY UNIVERSITY BY US.")
    
    add_p(doc, "", space_before=40)
    add_p(doc, "SAJITH AR (211423247183)\nSARAN R (211423247189)", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=30, space_after=12)
    doc.add_page_break()

    # ==========================================
    # 4. ACKNOWLEDGEMENT
    # ==========================================
    add_p(doc, "ACKNOWLEDGEMENT", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=12)
    
    def add_ack_p(doc, text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.30
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    ack_paras = [
        "We express our deep gratitude to our respected Secretary and Correspondent, Dr. P. CHINNADURAI, M.A., Ph.D., for his kind words, visionary leadership, and enthusiastic motivation, which inspired us immensely throughout the development of this project.",
        "We would like to extend our heartfelt thanks to our Directors, Tmt. C. VIJAYARAJESWARI, Dr. C. SAKTHIKUMAR, M.E., Ph.D., and Tmt. SARANYASREE SAKTHIKUMAR, B.E., M.B.A., for providing us with state-of-the-art computational infrastructure, laboratories, and an environment highly conducive to innovation.",
        "We express our sincere gratitude to our Principal, Dr. K. MANI, M.E., Ph.D., for his continuous encouragement, academic guidance, and invaluable support throughout our degree program.",
        "We also express our profound gratitude to our Chief Academic Officer, Dr. S. PRASANNA DEVI, M.E., Ph.D., PGDHRM, PDF, for her inspiring mentorship, timely guidance, and academic leadership.",
        "We convey our deepest gratitude to our Head of the Department of Artificial Intelligence and Machine Learning, Dr. G. SURESH, M.Tech., Ph.D., Professor, for his steadfast support, constructive technical reviews, and encouragement to pursue high-impact, socially relevant research.",
        "We would also like to express our sincere appreciation to our Project Coordinator, Dr. B. PRABHA, M.Tech., Ph.D., Assistant Professor Grade 1, for her rigorous reviews, valuable technical feedback, and seamless coordination of our project milestones.",
        "We would like to express our heartfelt indebtedness to our Supervisor, Mrs. UDHAYA J.M., M.E., Assistant Professor, for her patient mentorship, insightful domain expertise in artificial intelligence and optimization, and constant encouragement throughout every phase of this research.",
        "Finally, we thank all the faculty members, technical staff, our loving parents, and our fellow classmates in the Department of Artificial Intelligence and Machine Learning who supported us directly and indirectly in completing this endeavor."
    ]

    for para in ack_paras:
        add_ack_p(doc, para)
    
    p_names = doc.add_paragraph()
    p_names.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_names.paragraph_format.space_before = Pt(28)
    p_names.paragraph_format.space_after = Pt(0)
    p_names.paragraph_format.line_spacing = 1.25
    r_n = p_names.add_run("SAJITH AR (211423247183)\nSARAN R (211423247189)")
    r_n.font.name = 'Times New Roman'
    r_n.font.size = Pt(12)
    r_n.font.bold = True
    doc.add_page_break()

    # ==========================================
    # 5. ABSTRACT
    # ==========================================
    add_p(doc, "ABSTRACT", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=10)
    
    def add_abs_p(doc, text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    abs_1 = (
        "During severe urban flooding events, rapid and equitable emergency resource distribution is vital to "
        "minimizing casualties and human suffering. Traditional disaster response relies predominantly on post-incident "
        "field surveys, manual spreadsheet accounting, and static navigation routing tools that fail to account for "
        "dynamic road submersion. This inevitably causes critical supply stockouts at relief shelters, cognitive "
        "overload for municipal dispatch coordinators, and vehicle stranding along flooded arterial corridors. To resolve "
        "these challenges, this project presents an integrated, software-driven crisis management and logistics platform "
        "combining spatio-temporal inundation dynamics with machine learning and mathematical optimization, demonstrated "
        "on Chennai's urban network."
    )

    abs_2 = (
        "The proposed framework deploys a dual machine learning pipeline. For hazard classification, a 180-tree "
        "Gradient Boosted Decision Tree (GBDT) model is trained on 9,131 daily records (2001–2025) from the NASA POWER "
        "meteorological reanalysis dataset for Chennai, achieving a verified ROC-AUC of 0.8712 and 80.36% recall on four "
        "years of untouched held-out data (2022–2025). Simultaneously, a multi-output regression ensemble combining "
        "64 Extra Trees, 64 Random Forest trees, and scaled Ridge regression predicts per-capita emergency drinking water, "
        "food rations, and medical supply demands across 55,000 scenario records, reducing prediction errors over municipal "
        "planning rules by 88.1% to 96.6% under zero-leakage Cyclone Michaung holdout conditions. To guarantee operational "
        "reliability under demand volatility, distribution-free split-conformal calibration ensures nominal 90% finite-sample "
        "prediction intervals. Google OR-Tools Linear Programming (GLOP) solves multi-commodity supply allocation in under 65ms, "
        "while dynamic NetworkX graph routing evaluates road impedance against physical vehicle submersion thresholds (0.30m for "
        "trucks, 0.50m for pedestrian rescue). The platform is deployed as a sub-second (<0.45s) Next.js 14 / FastAPI GIS console "
        "supporting automated bilingual Tamil/English advisories and an offline 200-ward GCC atlas."
    )

    add_abs_p(doc, abs_1)
    add_abs_p(doc, abs_2)
    doc.add_page_break()

    # ==========================================
    # 6. TABLE OF CONTENTS
    # ==========================================
    add_p(doc, "TABLE OF CONTENTS", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18)
    
    toc_data = [
        ["ABSTRACT", "", "v"],
        ["TABLE OF CONTENTS", "", "vi"],
        ["LIST OF TABLES", "", "ix"],
        ["LIST OF FIGURES", "", "xi"],
        ["LIST OF SYMBOLS & ABBREVIATIONS", "", "xiii"],
        ["1.", "INTRODUCTION", "1"],
        ["1.1", "Background of Flood Disaster & Crisis Logistics", "1"],
        ["1.2", "Problem Statement", "2"],
        ["1.3", "Motivation of the Project", "3"],
        ["1.4", "Objectives of the Project", "4"],
        ["1.5", "Scope of the Project", "5"],
        ["2.", "LITERATURE SURVEY", "6"],
        ["2.1", "Overview of Urban Flood Modeling and Disaster Logistics", "6"],
        ["2.2", "Machine Learning for Spatio-Temporal Inundation & Precipitation", "6"],
        ["2.3", "Mathematical Optimization in Emergency Relief Distribution", "7"],
        ["2.4", "Dynamic Routing and Inundation Graph Traversal Algorithms", "8"],
        ["2.5", "Limitations of Existing Systems and Gap Analysis", "9"],
        ["3.", "PROPOSED SYSTEM", "12"],
        ["3.1", "Existing System Overview & Operational Bottlenecks", "12"],
        ["3.2", "Implementation Environment", "12"],
        ["3.3", "Proposed Methodology", "14"],
        ["3.4", "System Architecture", "15"],
        ["3.5", "Module Design & End-to-End Pipeline", "17"],
        ["4.", "SYSTEM IMPLEMENTATION", "20"],
        ["4.1", "Module Description", "20"],
        ["4.2", "Data Ingestion, Preprocessing & Spatio-Temporal Feature Engineering", "20"],
        ["4.3", "Machine Learning Formulation & Model Training", "23"],
        ["4.4", "Distribution-Free Split-Conformal Uncertainty Calibration", "26"],
        ["4.5", "Multi-Commodity Supply Allocation via Linear Programming", "27"],
        ["4.6", "Hazard-Aware Evacuation Routing via Dynamic Graph Traversal", "29"],
        ["4.7", "Geospatial Dashboard, Bilingual Bulletins & Automated Reporting", "30"],
        ["5.", "RESULTS AND DISCUSSION", "31"],
        ["5.1", "Weather Classification Performance", "31"],
        ["5.2", "Emergency Demand Prediction & Error Reduction Analysis", "33"],
        ["5.3", "Conformal Prediction Coverage & Empirical Calibration Validation", "34"],
        ["5.4", "Cross-Validation & Generalization Stability", "34"],
        ["5.5", "Computational Latency & Scalability Benchmarks", "36"],
        ["5.6", "Discussion on Observed Limitations & Edge Constraints", "37"],
        ["5.7", "Summary of Experimental Results", "38"],
        ["6.", "CONCLUSION AND FUTURE WORK", "39"],
        ["6.1", "Conclusion", "39"],
        ["6.2", "Societal & Environmental Impact", "40"],
        ["6.3", "Future Enhancements", "40"],
        ["", "APPENDIX", "42"],
        ["A.1", "Sustainable Development Goals (SDGs)", "43"],
        ["A.2", "System Source Code Listings", "45"],
        ["A.3", "System Execution Screenshots", "61"],
        ["", "REFERENCES", "65"]
    ]
    
    toc_table = doc.add_table(rows=1, cols=3)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc_table.autofit = False
    col_w = [Inches(1.0), Inches(4.5), Inches(0.9)]
    
    # Header
    hdr = toc_table.rows[0].cells
    hdr[0].width = col_w[0]
    hdr[1].width = col_w[1]
    hdr[2].width = col_w[2]
    hdr[0].paragraphs[0].add_run("CHAPTER NO.").font.bold = True
    hdr[1].paragraphs[0].add_run("TITLE").font.bold = True
    hdr[2].paragraphs[0].add_run("PAGE NO.").font.bold = True
    for c in hdr:
        p = c.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(11)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        
    for row_item in toc_data:
        row = toc_table.add_row()
        cells = row.cells
        cells[0].width = col_w[0]
        cells[1].width = col_w[1]
        cells[2].width = col_w[2]
        
        is_major = row_item[0] in ["1.", "2.", "3.", "4.", "5.", "6.", ""] or row_item[0] in ["ABSTRACT", "LIST OF TABLES", "LIST OF FIGURES", "LIST OF SYMBOLS & ABBREVIATIONS"]
        
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(row_item[0])
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        r0.font.bold = is_major
        
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(row_item[1])
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        r1.font.bold = is_major
        
        p2 = cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(row_item[2])
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10.5)
        r2.font.bold = is_major
        
    doc.add_page_break()

    # ==========================================
    # 7. LIST OF TABLES
    # ==========================================
    add_p(doc, "LIST OF TABLES", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18)
    
    tables_list = [
        ["Table 2.5.1", "Comparative Analysis of Literature Survey on Disaster Management Systems", "9"],
        ["Table 3.2.1", "Implementation Environment Specifications (Hardware, Software, Libraries)", "13"],
        ["Table 4.2.1", "NASA POWER Meteorological Feature Inventory & Scenario Attributes", "21"],
        ["Table 4.3.1", "Hyperparameter Specifications for 180-Tree GBDT Weather Classifier", "24"],
        ["Table 4.3.2", "Ensemble Architecture Parameters for Relief Demand Prediction", "26"],
        ["Table 4.4.1", "Split-Conformal Nonconformity Quantiles and Empirical Coverage Validation", "27"],
        ["Table 4.5.1", "Multi-Commodity Supply Optimization Parameters and Linear Constraints", "28"],
        ["Table 4.6.1", "Road Inundation Traversal Thresholds by Vehicle Classification", "29"],
        ["Table 5.1.1", "Next-Day Heavy Rainfall Hazard Classification Performance Metrics", "31"],
        ["Table 5.1.2", "Confusion Matrix Contingency Table for Rainfall Hazard Classification", "32"],
        ["Table 5.2.1", "Quantitative Demand Estimation Comparison (Proposed ML vs Municipal Baseline)", "33"],
        ["Table 5.4.1", "5-Fold Stratified Cross-Validation Stability Across Folds", "35"],
        ["Table 5.5.1", "Operational Latency and Computational Benchmarks across System Modules", "36"],
        ["Table 5.7.1", "Comprehensive Summary of Experimental Benchmark Results", "38"]
    ]
    
    t_table = doc.add_table(rows=1, cols=3)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_table.autofit = False
    col_w_t = [Inches(1.2), Inches(4.3), Inches(0.9)]
    
    hdr = t_table.rows[0].cells
    hdr[0].width = col_w_t[0]
    hdr[1].width = col_w_t[1]
    hdr[2].width = col_w_t[2]
    hdr[0].paragraphs[0].add_run("TABLE NO.").font.bold = True
    hdr[1].paragraphs[0].add_run("TITLE").font.bold = True
    hdr[2].paragraphs[0].add_run("PAGE NO.").font.bold = True
    for c in hdr:
        p = c.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(11)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        
    for item in tables_list:
        row = t_table.add_row()
        cells = row.cells
        cells[0].width = col_w_t[0]
        cells[1].width = col_w_t[1]
        cells[2].width = col_w_t[2]
        
        for c_idx, text in enumerate(item):
            p = cells[c_idx].paragraphs[0]
            if c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            if c_idx == 0:
                run.font.bold = True
                
    doc.add_page_break()

    # ==========================================
    # 8. LIST OF FIGURES
    # ==========================================
    add_p(doc, "LIST OF FIGURES", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18)
    
    figures_list = [
        ["Figure 3.4.1", "Layered 7-Tier System Architecture Diagram", "15"],
        ["Figure 3.5.1", "End-to-End Operational Execution Pipeline", "17"],
        ["Figure 4.2.1", "NASA POWER Meteorological Reanalysis & Scenario Dataset Distributions", "23"],
        ["Figure 4.3.1", "180-Tree GBDT Weather Classifier ROC and PR Curves", "25"],
        ["Figure 4.4.1", "Split-Conformal 90% Prediction Interval Calibration Across Relief Commodities", "27"],
        ["Figure 5.1.1", "Confusion Matrix for Next-Day Heavy Rainfall Hazard Classification", "33"],
        ["Figure 5.2.1", "Resource Demand Error Reduction: Machine Learning Ensemble vs Municipal Baseline", "34"],
        ["Figure 5.4.1", "5-Fold Stratified Cross-Validation ROC Curves Demonstrating Generalization Stability", "35"],
        ["Figure 5.5.1", "Supply Shortage Reduction and Latency Benchmarks", "37"],
        ["Figure A.3.1", "Full Executive Operations Dashboard with Real-Time Flood Inundation Overlay", "61"],
        ["Figure A.3.2", "Hazard-Aware Evacuation Routing Inspector with Road Submersion Verification", "61"],
        ["Figure A.3.3", "200 Greater Chennai Corporation Ward Vulnerability and Logistics Atlas", "62"],
        ["Figure A.3.4", "Multi-Disaster Scenario Comparison Engine (2015 vs 2023 vs Cyclone Michaung)", "62"],
        ["Figure A.3.5", "Historical Weather Replay & Precipitation Telemetry Ingestion", "63"],
        ["Figure A.3.6", "Multi-Modal Rescue Fleet Allocator (High-Clearance Trucks, Boats, Aerial Drops)", "63"],
        ["Figure A.3.7", "Automated Bilingual Tamil/English Disaster Advisory Dispatch Console", "64"],
        ["Figure A.3.8", "Responsive Mobile Field Dispatch Interface for Municipal First Responders", "64"]
    ]
    
    f_table = doc.add_table(rows=1, cols=3)
    f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_table.autofit = False
    col_w_f = [Inches(1.2), Inches(4.3), Inches(0.9)]
    
    hdr = f_table.rows[0].cells
    hdr[0].width = col_w_f[0]
    hdr[1].width = col_w_f[1]
    hdr[2].width = col_w_f[2]
    hdr[0].paragraphs[0].add_run("FIGURE NO.").font.bold = True
    hdr[1].paragraphs[0].add_run("TITLE").font.bold = True
    hdr[2].paragraphs[0].add_run("PAGE NO.").font.bold = True
    for c in hdr:
        p = c.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(11)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        
    for item in figures_list:
        row = f_table.add_row()
        cells = row.cells
        cells[0].width = col_w_f[0]
        cells[1].width = col_w_f[1]
        cells[2].width = col_w_f[2]
        
        for c_idx, text in enumerate(item):
            p = cells[c_idx].paragraphs[0]
            if c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            if c_idx == 0:
                run.font.bold = True
                
    doc.add_page_break()

    # ==========================================
    # 9. LIST OF SYMBOLS & ABBREVIATIONS
    # ==========================================
    add_p(doc, "LIST OF SYMBOLS & ABBREVIATIONS", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18)
    
    symbols_data = [
        ["AI", "Artificial Intelligence"],
        ["ML", "Machine Learning"],
        ["GBDT", "Gradient Boosted Decision Trees"],
        ["ExtraTrees", "Extremely Randomized Trees"],
        ["OR-Tools", "Google Operations Research Tools"],
        ["GLOP", "Google Linear Optimization Package"],
        ["GIS", "Geographic Information System"],
        ["GCC", "Greater Chennai Corporation"],
        ["IMD", "India Meteorological Department"],
        ["NASA POWER", "Prediction of Worldwide Energy Resources (NASA)"],
        ["ROC-AUC", "Receiver Operating Characteristic – Area Under Curve"],
        ["PR-AUC", "Precision-Recall – Area Under Curve"],
        ["MAE", "Mean Absolute Error"],
        ["RMSE", "Root Mean Square Error"],
        ["CV", "Cross-Validation"],
        ["LP", "Linear Programming"],
        ["MIP", "Mixed-Integer Programming"],
        ["VRPTW", "Vehicle Routing Problem with Time Windows"],
        ["CPLEX", "IBM ILOG CPLEX Optimization Studio"],
        ["REST", "Representational State Transfer"],
        ["API", "Application Programming Interface"],
        ["JSON", "JavaScript Object Notation"],
        ["GeoJSON", "Geographic JavaScript Object Notation"],
        ["SDG", "Sustainable Development Goal (United Nations)"],
        ["D_i^k", "Demand of commodity k at shelter or distribution node i"],
        ["S_j^k", "Available stock of commodity k at warehouse or depot j"],
        ["x_{ij}^k", "Quantity of commodity k allocated from depot j to shelter i"],
        ["c_{ij}", "Unit transportation cost or distance impedance between depot j and shelter i"],
        ["s_i^k", "Unmet demand slack variable for commodity k at shelter i"],
        ["w_i", "Vulnerability priority weight for shelter or ward node i"],
        ["d_{uv}", "Flood water depth along road edge (u, v)"],
        ["\\theta_{veh}", "Critical physical flood depth threshold for vehicle category veh"],
        ["\\alpha", "Significance level for conformal prediction coverage (1 - \\alpha = 0.90)"],
        ["\\hat{q}", "Conformal empirical nonconformity quantile at level (1 - \\alpha)"]
    ]
    
    s_table = doc.add_table(rows=1, cols=2)
    s_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_table.autofit = False
    col_w_s = [Inches(1.8), Inches(4.6)]
    
    hdr = s_table.rows[0].cells
    hdr[0].width = col_w_s[0]
    hdr[1].width = col_w_s[1]
    hdr[0].paragraphs[0].add_run("ABBREVIATION / SYMBOL").font.bold = True
    hdr[1].paragraphs[0].add_run("EXPANSION / DEFINITION").font.bold = True
    for c in hdr:
        p = c.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(11)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        
    for item in symbols_data:
        row = s_table.add_row()
        cells = row.cells
        cells[0].width = col_w_s[0]
        cells[1].width = col_w_s[1]
        
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(item[0])
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(item[1])
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        
    doc.add_page_break()
